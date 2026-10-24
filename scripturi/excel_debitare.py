"""Liste de debitare + gaurire + feronerie pentru dulapul din hol (ap. 184), pentru Debimax si Mambricolaj."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTDIR = sys.argv[1] if len(sys.argv) > 1 else "."

W18 = "Egger W1000 ST9 Premium White"
U18 = "Egger U702 ST9 Cashmere Grey"
W16 = "Egger W1000 ST9 Premium White"
HDF = "HDF alb"

# ---------------------------------------------------------------- piese (mm)
# cod, denumire, material, grosime, L, l, buc, cant L1, L2, l1, l2, gaurire, observatii
# L1 = latura lunga din fata (vizibila), L2 = latura lunga opusa, l1 / l2 = laturile scurte
P = [
    ("C01", "Lateral stânga", W18, 18, 2380, 580, 1, "2 mm", "-", "0,8 mm", "0,8 mm", "DA – vezi foaia Găurire", "L1 = muchia din față"),
    ("C02", "Lateral oblic (dreapta)", W18, 18, 2380, 599, 1, "2 mm", "-", "0,8 mm", "0,8 mm", "DA – vezi foaia Găurire", "Se montează la 14,5° față de front"),
    ("C03", "Bază", W18, 18, 2417, 580, 1, "2 mm", "-", "-", "-", "DA – vezi foaia Găurire", "TRAPEZ: capătul drept tăiat oblic, latura spate = 2267 mm"),
    ("C04", "Blat sus", W18, 18, 2417, 580, 1, "2 mm", "-", "-", "-", "DA – vezi foaia Găurire", "TRAPEZ: capătul drept tăiat oblic, latura spate = 2267 mm"),
    ("C05", "Despărțitor central", W18, 18, 2344, 577, 1, "2 mm", "-", "-", "-", "DA – vezi foaia Găurire", ""),
    ("C06", "Despărțitor colț", W18, 18, 2344, 577, 1, "2 mm", "-", "-", "-", "DA – vezi foaia Găurire", ""),
    ("C07", "Poliță fixă modul stânga", W18, 18, 1017, 557, 2, "2 mm", "-", "-", "-", "DA – vezi foaia Găurire", "1 deasupra sertarului + 1 sus"),
    ("C08", "Poliță fixă modul dreapta", W18, 18, 1027, 557, 2, "2 mm", "-", "-", "-", "DA – vezi foaia Găurire", "1 deasupra sertarului + 1 sus"),
    ("C09", "Raft colț", W18, 18, 577, 337, 4, "-", "-", "2 mm", "-", "NU", "TRAPEZ: l1 = față 337 mm, latura spate = 188 mm (tăiere oblică)"),
    ("C10", "Soclu (retras)", W18, 18, 2396, 80, 1, "-", "0,8 mm", "-", "-", "NU", "L2 = muchia de jos"),
    ("F01", "Ușă jos (suport oglindă)", U18, 18, 1699, 519, 4, "2 mm", "2 mm", "2 mm", "2 mm", "DA – balamale Ø35", "Oglinda 4 mm se lipește pe față după montaj"),
    ("F02", "Ușă sus (suport oglindă)", U18, 18, 377, 519, 4, "2 mm", "2 mm", "2 mm", "2 mm", "DA – balamale Ø35", ""),
    ("F03", "Front sertar stânga", U18, 18, 1041, 366, 1, "2 mm", "2 mm", "2 mm", "2 mm", "NU", "Push-to-open, fără mâner"),
    ("F04", "Front sertar dreapta", U18, 18, 1042, 366, 1, "2 mm", "2 mm", "2 mm", "2 mm", "NU", "Push-to-open, fără mâner"),
    ("F05", "Plintă colț", U18, 18, 363, 88, 1, "2 mm", "-", "2 mm", "2 mm", "NU", "L1 = muchia de sus"),
    ("F06", "Mască tavan", U18, 18, 2453, 60, 1, "2 mm", "-", "-", "2 mm", "NU", "L1 = jos; L2 se ajustează pe tavan la montaj"),
    ("S01", "Lateral sertar", W16, 16, 490, 200, 4, "0,8 mm", "-", "-", "-", "DA – vezi foaia Găurire", "L1 = muchia de sus"),
    ("S02", "Față interioară + spate sertar stânga", W16, 16, 975, 171, 2, "0,8 mm", "-", "-", "-", "DA – vezi foaia Găurire", "Spate: + găuri cârlig MOVENTO"),
    ("S03", "Față interioară + spate sertar dreapta", W16, 16, 985, 171, 2, "0,8 mm", "-", "-", "-", "DA – vezi foaia Găurire", "Spate: + găuri cârlig MOVENTO"),
    ("S04", "Fund sertar stânga", W16, 16, 975, 490, 1, "-", "-", "-", "-", "NU", ""),
    ("S05", "Fund sertar dreapta", W16, 16, 985, 490, 1, "-", "-", "-", "-", "NU", ""),
    ("H01", "Spate HDF stânga", HDF, 3, 2376, 1044, 1, "-", "-", "-", "-", "NU", ""),
    ("H02", "Spate HDF dreapta", HDF, 3, 2376, 1256, 1, "-", "-", "-", "-", "NU", "Se îmbină pe despărțitorul central"),
]

# ---------------------------------------------------------------- gaurire
# piesa, fata, tip, diametru, adancime, X, Y, buc, reper
LOWER_A = [100, 600, 1099, 1599]          # usi 1, 2, 4 (mm de la muchia de jos a usii)
LOWER_B = [140, 640, 1139, 1559]          # usa 3 (decalat, ca placutele sa nu se intalneasca in despartitor)
UPPER_A = [80, 297]
UPPER_B = [120, 257]
DOOR_BOT = 379.5                          # muchia de jos a usilor de jos, de la pardoseala
UP_BOT = 2081.5


def plates(lst, base, off):
    return [round(base + y - off, 1) for y in lst]


G = []
ref_v = "X = de la muchia din față, Y = de la muchia de jos a piesei"
# C01 lateral stanga (jos la 80 mm de la pardoseala)
for y in (9, 2371):
    for x in (50, 290, 530):
        G.append(("C01", "interior", "Confirmat (față)", "Ø7 străpuns", "-", x, y, 1, ref_v))
for y in (289, 2009):
    for x in (73, 527):
        G.append(("C01", "interior", "Confirmat poliță (față)", "Ø7 străpuns", "-", x, y, 1, ref_v))
for y in plates(LOWER_A, DOOR_BOT, 80) + plates(UPPER_A, UP_BOT, 80):
    for dy in (-16, 16):
        G.append(("C01", "interior", "Plăcuță balama Blum 173L6100", "Ø5", "11", 37, round(y + dy, 1), 1, ref_v))
# C02 lateral oblic
for y in (9, 2371):
    for x in (50, 300, 549):
        G.append(("C02", "interior", "Confirmat (față)", "Ø7 străpuns", "-", x, y, 1, ref_v))
for y in (413, 813, 1213, 1613):
    for x in (60, 540):
        G.append(("C02", "interior", "Pin raft", "Ø5", "10", x, y, 1, ref_v))
# C05 despartitor central (jos la 98 mm)
for y in (271, 1991):
    for x in (73, 527):
        G.append(("C05", "ambele", "Confirmat poliță stânga", "Ø7 străpuns", "-", x, y, 1, ref_v))
    for x in (123, 477):
        G.append(("C05", "ambele", "Confirmat poliță dreapta", "Ø7 străpuns", "-", x, y, 1, ref_v))
for y in plates(LOWER_A, DOOR_BOT, 98) + plates(UPPER_A, UP_BOT, 98):
    for dy in (-16, 16):
        G.append(("C05", "stânga (modul A)", "Plăcuță balama ușa 2", "Ø5", "11", 37, round(y + dy, 1), 1, ref_v))
for y in plates(LOWER_B, DOOR_BOT, 98) + plates(UPPER_B, UP_BOT, 98):
    for dy in (-16, 16):
        G.append(("C05", "dreapta (modul B)", "Plăcuță balama ușa 3", "Ø5", "11", 37, round(y + dy, 1), 1, ref_v))
for x in (100, 480):
    G.append(("C05", "cant jos", "Cep aliniere", "Ø8", "25", x, "-", 1, "X = de la muchia din față"))
# C06 despartitor colt
for y in (271, 1991):
    for x in (73, 527):
        G.append(("C06", "ambele", "Confirmat poliță dreapta", "Ø7 străpuns", "-", x, y, 1, ref_v))
for y in plates(LOWER_A, DOOR_BOT, 98) + plates(UPPER_A, UP_BOT, 98):
    for dy in (-16, 16):
        G.append(("C06", "stânga (modul B)", "Plăcuță balama ușa 4", "Ø5", "11", 37, round(y + dy, 1), 1, ref_v))
for y in (395, 795, 1195, 1595):
    for x in (60, 517):
        G.append(("C06", "dreapta (colț)", "Pin raft", "Ø5", "10", x, y, 1, ref_v))
for x in (100, 480):
    G.append(("C06", "cant jos", "Cep aliniere", "Ø8", "25", x, "-", 1, "X = de la muchia din față"))
# C03 baza / C04 blat
ref_h = "X = de la capătul stâng, Y = de la muchia din față"
for x in (1026, 2071):
    for y in (100, 480):
        G.append(("C03", "sus", "Cep aliniere despărțitor", "Ø8", "12", x, y, 1, ref_h))
for pc in ("C03", "C04"):
    for y in (50, 290, 530):
        G.append((pc, "cant stânga", "Confirmat (cant)", "Ø5", "45", "-", y, 1, "Y = de la muchia din față"))
    G.append((pc, "cant oblic dreapta", "Confirmat (cant)", "Ø5", "45", "-", "50 / 300 / 549 (pe oblic)", 3, "LA MONTAJ, cu șablon – cant oblic"))
# C07 / C08 polite
for pc, left, right in (("C07", (50, 504), (50, 504)), ("C08", (100, 454), (50, 504))):
    for y in left:
        G.append((pc, "cant stânga", "Confirmat (cant)", "Ø5", "45", "-", y, 2, "Y = de la muchia din față; buc = pe ambele polițe"))
    for y in right:
        G.append((pc, "cant dreapta", "Confirmat (cant)", "Ø5", "45", "-", y, 2, "Y = de la muchia din față; buc = pe ambele polițe"))
# usi
ref_d = "X = de la muchia cu balamale, Y = de la muchia de jos"
for door, lst in (("F01 – ușa 1 (balama stânga)", LOWER_A), ("F01 – ușa 2 (balama dreapta)", LOWER_A),
                  ("F01 – ușa 3 (balama stânga)", LOWER_B), ("F01 – ușa 4 (balama dreapta)", LOWER_A),
                  ("F02 – ușa 1 (balama stânga)", UPPER_A), ("F02 – ușa 2 (balama dreapta)", UPPER_A),
                  ("F02 – ușa 3 (balama stânga)", UPPER_B), ("F02 – ușa 4 (balama dreapta)", UPPER_A)):
    for y in lst:
        G.append((door, "spate", "Locaș balama Blum CLIP top", "Ø35", "13", 21.5, y, 1, ref_d))
        G.append((door, "spate", "Șuruburi balama (45/9,5)", "Ø2,5", "10", 31, "%s ± 22,5" % y, 2, ref_d))
# sertare
for y in (69, 160):
    for x in (8, 482):
        G.append(("S01", "interior", "Confirmat (față)", "Ø7 străpuns", "-", x, y, 4, "X = de la capătul din față, Y = de la muchia de jos; buc = pe toate 4 lateralele"))
for pc in ("S02", "S03"):
    for y in (40, 131):
        G.append((pc, "cant ambele capete", "Confirmat (cant)", "Ø5", "45", "-", y, 4, "Y = de la muchia de jos; buc = 2 capete × 2 piese"))
    G.append((pc, "spate – jos", "Cârlig posterior MOVENTO", "Ø6", "10", "conform fișă Blum", "-", 2, "DE CONFIRMAT cu fișa tehnică MOVENTO"))

# ---------------------------------------------------------------- feronerie
FER = [
    ("Blum", "70T3550", "Balama CLIP top 110°, fără arc, aplicată", 6, "ușa 1 (4 jos + 2 sus); fără arc = compatibil TIP-ON"),
    ("Blum", "70T3650*", "Balama CLIP top 110°, fără arc, semiaplicată", 18, "ușile 2, 3, 4"),
    ("Blum", "173L6100", "Plăcuță montaj 0 mm, cruce", 24, ""),
    ("Blum", "956.1004", "TIP-ON pentru uși, varianta scurtă", 8, "+ adaptor drept pentru montaj în corp (cod la dealer)"),
    ("Blum", "760H5000S", "Glisiere MOVENTO BLUMOTION S, NL 500, 40 kg", "2 perechi", "extragere totală"),
    ("Blum", "*", "TIP-ON BLUMOTION pentru MOVENTO", "2 seturi", "push-to-open pentru sertare"),
    ("Blum", "*", "Sincronizare TIP-ON BLUMOTION, multi-part", 2, "se taie la 750 mm (stânga) și 760 mm (dreapta)"),
    ("Blum", "*", "Cleme de cuplare front MOVENTO (locking devices)", "2 perechi", ""),
    ("Diverse", "", "Confirmat 7 × 50 + cheie hexagonală", 50, "44 necesare + rezervă"),
    ("Diverse", "", "Capace confirmat albe", 50, ""),
    ("Diverse", "", "Cepuri lemn Ø8 × 35", 4, "despărțitoare pe bază"),
    ("Diverse", "", "Colțar metalic mobilă 40 × 40", 10, "6 despărțitoare + 4 mască tavan"),
    ("Diverse", "", "Colțar anti-răsturnare + diblu 8 × 60 + șurub 5 × 60", 3, "prindere în perete – OBLIGATORIU"),
    ("Diverse", "", "Picior reglabil soclu 80–100 mm", 10, ""),
    ("Diverse", "", "Clemă soclu", 5, ""),
    ("Diverse", "", "Pin raft Ø5", 16, "rafturi colț"),
    ("Diverse", "", "Bară ovală aluminiu 30 × 15", "2 (1015 + 1025 mm)", "+ 4 suporți capăt"),
    ("Diverse", "", "Șurub PAL 4 × 16", 250, "balamale, plăcuțe, TIP-ON, glisiere, colțare, picioare"),
    ("Diverse", "", "Șurub PAL 4 × 30", 20, "fronturi sertare, funduri sertare"),
    ("Diverse", "", "Șurub PAL 4 × 40", 12, "lateral sertar în fund"),
    ("Diverse", "", "Șurub 3,5 × 16 sau capse", 80, "spate HDF"),
    ("Sticlărie", "", "Oglindă argintie 4 mm cu folie de siguranță 1695 × 515", 4, "ușile de jos"),
    ("Sticlărie", "", "Oglindă argintie 4 mm cu folie de siguranță 373 × 515", 4, "ușile de sus"),
    ("Diverse", "", "Adeziv oglinzi (silicon neutru)", 2, "tuburi"),
    ("LED", "", "Profil aluminiu aplicat + difuzor opal, 1 m", 2, "sub polița de sus"),
    ("LED", "", "Bandă LED 24 V, 4000 K, CRI > 90, ~10 W/m", "2 m", ""),
    ("LED", "", "Senzor de ușă IR 24 V, autoadeziv (24led, 5–24 V DC, 5 A)", 4, "pe corp, lângă balamaua de sus a celor 4 uși mari; câte 2 în paralel pe modul"),
    ("LED", "", "Sursă 24 V, 30–60 W", 1, "în compartimentul de sus; doză 230 V la 240 cm"),
]

NOTE = [
    "Proiect: Apartament 184, Bloc L14 – dulap hol, construit pe tot peretele, cu capătul drept oblic (14,5°).",
    "Toate dimensiunile sunt în mm, finite, CU cant inclus.",
    "Decorurile W1000 și U702 sunt unicolore: rotirea pieselor la optimizare este PERMISĂ.",
    "Cant: L1 = latura lungă din față (vizibilă), L2 = latura lungă opusă, l1 / l2 = laturile scurte.",
    "C03 / C04 (bază, blat): trapez – capătul din dreapta se taie oblic; latura din față 2417, latura din spate 2267.",
    "C09 (rafturi colț): trapez – față 337, spate 188, adâncime 577.",
    "Găurire: Ø35 = locaș balama, Ø5 × 11 = plăcuță balama (linia de 37 mm), Ø7 străpuns + Ø5 × 45 în cant = confirmat, Ø5 × 10 = pin raft, Ø8 = cep.",
    "Găurile în cant marcate „LA MONTAJ” (cant oblic) le facem noi la montaj, cu șablon.",
    "Dacă atelierul nu face găuri în cant, le facem noi cu șablon – vă rugăm să ne spuneți.",
    "Înainte de debitare: înălțimea finită a tavanului în hol trebuie să fie ≥ 2480 mm (dulapul are 2460 mm + mască ajustabilă).",
    "Oglinzile se comandă separat, la sticlărie (nu fac parte din debitare).",
    "Codurile Blum marcate cu * se confirmă la dealer / în configuratorul Blum.",
]

THIN = Side(style="thin", color="C8CCD2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="005E99")
ZEBRA = PatternFill("solid", fgColor="F2F5F8")


def sheet(ws, title, headers, rows, widths):
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13, color="005E99")
    ws.append([])
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    for i, r in enumerate(rows):
        ws.append(list(r))
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=4 + i, column=c)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if i % 2:
                cell.fill = ZEBRA
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = "A4"


def cant_meters():
    tot = {}
    for cod, den, mat, gr, L, l, n, c1, c2, c3, c4, _, _ in P:
        for c, length in ((c1, L), (c2, L), (c3, l), (c4, l)):
            if c != "-":
                key = "%s – %s" % (mat.replace("Egger ", "").split(" Premium")[0].split(" Cashmere")[0], c)
                tot[key] = tot.get(key, 0) + length * n / 1000
    return tot


def build(shop, extra_notes):
    wb = Workbook()
    ws = wb.active
    ws.title = "Debitare"
    rows = [(i + 1, p[0], p[1], p[2], p[3], p[4], p[5], p[6], "DA" if p[2] != HDF else "-", p[7], p[8], p[9], p[10], p[11], p[12])
            for i, p in enumerate(P)]
    sheet(ws, "Listă debitare – dulap hol, ap. 184 – pentru %s" % shop,
          ["Nr.", "Cod", "Denumire", "Material (decor)", "Gros. (mm)", "Lungime L (mm)", "Lățime l (mm)", "Buc.",
           "Rotire permisă", "Cant L1", "Cant L2", "Cant l1", "Cant l2", "Găurire CNC", "Observații"],
          rows, [5, 6, 30, 30, 7, 10, 10, 6, 9, 9, 9, 9, 9, 20, 46])
    r = ws.max_row + 2
    ws.cell(row=r, column=3, value="Necesar cant estimat (+10 %)").font = Font(bold=True)
    for k, v in sorted(cant_meters().items()):
        r += 1
        ws.cell(row=r, column=3, value=k)
        ws.cell(row=r, column=4, value="%.1f m" % (v * 1.1))
    r += 2
    ws.cell(row=r, column=3, value="Plăci estimate (2800 × 2070)").font = Font(bold=True)
    for k, v in (("W1000 ST9, 18 mm", "3 plăci (din a 3-a rămâne ~50 %)"), ("U702 ST9, 18 mm", "2 plăci (din a 2-a rămâne ~60 %)"),
                 ("W1000 ST9, 16 mm", "1 placă (~35 % folosit)"), ("HDF alb 3 mm", "2 plăci")):
        r += 1
        ws.cell(row=r, column=3, value=k)
        ws.cell(row=r, column=4, value=v)

    ws2 = wb.create_sheet("Găurire")
    sheet(ws2, "Plan de găurire – dulap hol, ap. 184 (mm)",
          ["Piesă", "Fața", "Tip", "Diametru", "Adâncime", "X", "Y", "Buc.", "Reper cote"],
          G, [30, 18, 32, 12, 10, 16, 22, 6, 52])

    ws3 = wb.create_sheet("Feronerie")
    sheet(ws3, "Feronerie și accesorii – dulap hol, ap. 184",
          ["Producător", "Cod", "Denumire", "Cantitate", "Observații"], FER, [12, 12, 52, 18, 52])

    ws4 = wb.create_sheet("Note atelier")
    ws4.append(["Note pentru %s" % shop])
    ws4["A1"].font = Font(bold=True, size=13, color="005E99")
    ws4.append([])
    for n in NOTE + extra_notes:
        ws4.append([n])
    ws4.column_dimensions["A"].width = 130
    fn = "%s\\Dulap_hol_ap184_%s.xlsx" % (OUTDIR, shop.replace(" ", "_"))
    wb.save(fn)
    return fn


print(build("Debimax", [
    "Debimax: vă rugăm să confirmați ce găuri puteți executa pe CNC (față / cant) și dacă tăiați oblic capetele pieselor C03, C04, C09.",
]))
print(build("Mambricolaj", [
    "Mambricolaj: găurire pentru balamale + sistem confirmat (ericsoane) conform foii „Găurire”.",
    "Dacă formularul vostru oficial este obligatoriu, datele din foaia „Debitare” se copiază în el coloană cu coloană.",
    "Verificăm în magazin disponibilitatea feroneriei Blum din foaia „Feronerie”.",
]))
print("piese:", len(P), "| randuri gaurire:", len(G))
