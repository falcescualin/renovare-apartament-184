"""Deviz R03: cea mai bună ofertă găsită pentru fiecare articol (preț verificat, magazin, link) + alternativa la un
magazin comod, etape de cheltuieli și plăți pe săptămâni, plus calendarul lucrărilor.
Pornește din R02 (tmp/deviz_R02.json). Scrie tmp/deviz.json (R03), tmp/timeline.json și Deviz_materiale_ap184_R03.xlsx.
Prețuri cu TVA, observate pe 05.10.2026. Ruleaza: python deviz_R03_oferte_etape.py"""
import json, shutil, datetime, collections
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / 'tmp/deviz.json'
BASE = ROOT / 'tmp/deviz_R02.json'
if not BASE.exists():
    shutil.copy(SRC, BASE)
D = json.loads(BASE.read_text(encoding='utf-8'))
assert D['revision'] == 'R02', D['revision']
ROWS = {r['id']: r for r in D['rows']}
V, E = 'verificat', 'estimat'
DT = '05.10.2026'


def bo(i, best, store, url, alt, alt_store, st=V, art=None, q=None, um=None, note=None):
    """cea mai bună ofertă (preț, magazin, link) și alternativa"""
    r = ROWS[i]
    r.update(pmin=round(best, 2), pmax=round(alt, 2), url=url, st=st, magazin=store, magazin_alt=alt_store,
             price_basis='Cea mai bună ofertă: %s; alternativă: %s (%s)' % (store, alt_store, DT))
    if art: r['art'] = art
    if q is not None: r['q'] = q
    if um: r['um'] = um
    if note: r['note'] = note


DD = 'https://www.dedeman.ro/ro/'
# ================================================================ 3. instalatie electrica
bo('M018', 378.99, 'Leroy Merlin 378,99 lei / rolă 100 m', 'https://www.leroymerlin.ro/produse/cabluri/625/cablu-electric-cyy-f-3-x-1.5-mm-100-m-gri/46894',
   484, 'Dedeman 484 lei', art='Cablu CYY-F 3 × 1,5 mm², o rolă de 100 m (necesar 86 m)', q=1, um='rolă')
bo('M019', 595 + 37 * 8.26, 'Hornbach 595 lei / rolă 100 m + 37 m la Micul Electrician (8,26 lei/m)',
   'https://www.hornbach.ro/p/cablu-cyy-f-3x2-5-mm-gri-inel-100m-manta-din-pvc-tip-st2-conform-sr-cei-60502/5800492/',
   725 + 37 * 7.25, 'Dedeman 725 lei / rolă + 37 m', art='Cablu CYY-F 3 × 2,5 mm²: rolă 100 m + 37 m (necesar 137 m)', q=1, um='lot')
bo('M020', 5 * 28.99 + 45 * 0.73 + 6 * 9.96, 'Dedeman: copex Total Plast 20 mm 28,99 lei / 50 m × 5; doze aparat 0,73 lei × 45; doze derivație 9,96 lei × 6',
   DD + 'copex-hdpe-total-plast-fara-halogen-d-20-mm-rola-50-m/p/1059633',
   250 * 2.89 + 45 * 0.75 + 6 * 9.96, 'copex Legrand 750 N 2,89 lei/m', q=1, um='lot')
bo('M021', (11 * 11.40 + 12 * 20.78) / 35, 'Micul Electrician: Asfora simplă 11,40 lei × 11; Dedeman: dublă 20,78 lei × 12 (35 posturi)',
   'https://miculelectrician.ro/priza-simpla-2pe-cu-cp-obturatoare-asfora-alb-16a-schneider',
   (11 * 14.90 + 12 * 20.78) / 35, 'Dedeman (simplă 14,90 lei)')
bo('M022', 121.60 / 8, 'Dedeman: Asfora simplu 14,24 lei × 6 + dublu 18,08 lei × 2', DD + 'intrerupator-dublu-schneider-electric-asfora-eph0300121-incastrat-rama-inclusa-alb/p/1036844',
   121.60 / 8, 'Dedeman')
bo('M023', 179 + 62.32 + 285 + 2 * 16.40 + 29.80, 'Tablou Noark 54 module 179 lei (Materiale Electrice Online) + separator iSW 62,32 lei (Dedeman) + SPD Easy9 285 lei (Hornbach) + bare N/PE 2 × 16,40 lei + pieptene 29,80 lei',
   'https://materialeelectriceonline.ro/tablouri-electrice/policarbonat/tablouri-aparente/2878-tablou-aparent-54-moduleip-40-3-randuri-3x18m-noark-pns-3x18t-106196.html',
   333 + 79.45 + 319.16 + 2 * 32.19 + 41.73, 'Tablou Gewiss 333 lei (Dedeman) + aparataj la prețul mai mare',
   art='Tablou 54 module aparent + separator 2P 40 A + descărcător SPD T2 1P+N + bare N/PE + pieptene')
bo('M024', 140.17, 'Electricon: ETI KZS-1M-UNI 002176024 tip A B16 30 mA, 140,17 lei (1 modul)', 'https://www.electricon.ro/siguran%C8%9Be-%C8%99i-protec%C8%9Bii-electrice/002176024.html',
   219, 'Vonmag: Schneider iCV40H A9DG4616 tip A, 219 lei', art='RCBO 1P+N 30 mA tip A, curba B (2 × B10 + 8 × B16), ETI KZS-1M-UNI sau Schneider iCV40H')
bo('M025', 994.30, 'Celon 994,30 lei', 'https://www.celon.ro/aparataje-modulare-1409/dispozitive-de-protectie-impotriva-arcului-electric-8239/acti9-afdd-intreruptor-automat-diferential-detectie-arc-electric-icv40n-1pn-curba-b-16a-6ka-30ma-tip-a-si-vigiarc-rcbo-1053962.html',
   1099, 'Vonmag 1.099 lei')
bo('M026', 86.50 + 20.99 + 50, 'Mathaus: Wago 221-413 50 buc 86,50 lei; eMAG etichete 20,99 lei; coliere ~50 lei',
   'https://mathaus.ro/p/conector-imbinare-cu-parghii-wago-221-413-30-a-3-x-4-mmp-50-bucati-set/000000000010514667',
   95 + 20.99 + 115, 'Dedeman Wago 95 lei')
# ================================================================ 5. LED: senzori de ușă (pe balamalele Blum nu există întrerupător)
bo('M039', 45.50, '24led: întrerupător de dulap 5–24 V DC 5 A, senzor IR autoadeziv, 45,50 lei', 'https://www.24led.ro/accesorii-banda-led/intrerupator-de-dulap-5-24v-dc-5a-usa.html',
   57.99, 'senzor IR mobilier 57,99 lei (classlights)', art='Senzori de ușă IR 24 V pentru LED (dulap hol 4, dressing 5)', q=9, um='buc',
   note='Senzorul se lipește pe corp lângă balamaua de sus și nu apasă pe ușă, deci merge cu TIP-ON. Alternativă: senzor pentru uși duble ~52 lei (LED Solution), câte unul pe modul.')
# ================================================================ 4. curenti slabi
# Cat6 și coaxialul merg în stea: fiecare priză are cablul ei până la distribuitor (R04: Cat6 4 legături 40,7 m, coax 3 legături 29,1 m)
bo('M030', 35 * 1.05 + 71.34 + 30, 'Dedeman: RG6 1,05 lei/m × 35 m (3 legături separate, 29,1 m + rezerve) + 3 prize TV Asfora 23,78 lei + distribuitor ~30 lei',
   DD + 'priza-tv-de-capat-asfora-eph3200121-rama-inclusa-alb/p/1036847',
   35 * 3.13 + 82.08 + 60, 'RG6 3,13 lei/m + prize 27,36 lei', art='Cablu coaxial RG6 35 m + 3 prize TV Asfora + distribuitor')
# ================================================================ 6. sanitare
bo('M040', 842.03, 'absulo.ro 842,03 lei', 'https://absulo.ro/p/roca-contesa-cada-dreptunghiulara-150x70-cm-alb-a2360k0000', 855.93, 'Dedeman 855,93 lei')
bo('M044', 209, 'Neakaisa 209 lei', 'https://www.neakaisa.ro/clapeta-dubla-actionare-alpin-alb-geberit-sigma01.html', 234, 'Dedeman 234 lei')
bo('M046', 2378.27, 'Dedeman 2.378,27 lei (cel mai mic preț găsit pentru A851476806)', DD + 'masca-baie-lavoar-roca-gap-new-60-a851476806-cu-sertare-alb-suspendat-60-5-x-46-0-x-53-7-cm/p/3044509',
   2760, 'absulo.ro 2.760 lei', note='Variantă mai ieftină, aproape identică: set Roca The Gap 60 A855997806 la German Quality, 2.199,94 lei.')
bo('M047', 269 + 60, 'Sanitino / Dedeman 269 lei + sifon ~60 lei', 'https://www.sanitino.ro/grohe-eurosmart-baterie-lavoar-monocomanda-s-cu-evacuare-crom-33265002',
   349 + 120, 'Dedeman cu ventil 349 lei + sifon')
bo('M049', 2 * 17.30 + 20 + 2 * 25, 'Dedeman: 2 robineți colțar Ferro 1/2 × 3/8 17,30 lei; robinet mașină ~20 lei; 2 robineți de izolare ~25 lei',
   DD + 'robinet-coltar-ferro-alama-cromata-filet-exterior-1/2-x-3/8-pn10/p/2024527',
   2 * 26.03 + 30 + 2 * 35, 'robineți cu cartuș ceramic', art='Robineți: 2 colțar lavoar, 1 mașină de spălat, 2 de izolare AR / AC',
   note='Calculat pe planul sanitar: lavoar AR + AC, racord mașină de spălat în vestibul, izolare la coloană.')
bo('M050', 5 * 9.75 + 16 * 0.76 + 4 * 0.76 + 7 * 7.96 + 8 * 4.71, 'Dedeman: țeavă PPR 20 PN20 9,75 lei / 4 m × 5; coturi și teuri 0,76 lei × 20; coturi cu filet 7,96 lei × 7; izolație 22 mm 4,71 lei × 8',
   DD + 'teava-ppr-20-ubp-4-m-x-20-mm-x-2-8-mm-cu-fibra-compozita-pn-20/p/2036515',
   5 * 12.50 + 20 * 1.0 + 7 * 9 + 8 * 5.94, 'țeavă PN20 3,4 mm 12,50 lei', art='Alimentare apă PPR 20 mm: ~10 m AR + ~6 m AC, fitinguri, izolație',
   note='Lungimi din modelul băii: coloană → rezervor 1,2 m, lavoar 1,8 m, baterie cadă 2,7 m, mașina de spălat ~2,5 m (AR), plus 20 % rezervă.')
bo('M051', 5 * 14.16 + 3 * 7 + 1 * 20 + 6 * 3.11 + 4 * 3 + 2 * 6 + 12 + 2 * 3 + 10,
   'Dedeman: PP Ø50 1 m 14,16 lei × 5; cot 50 87° 3,11 lei × 6; Ø40 × 3, Ø110 × 1, coturi 45°, ramificații, reducții estimate',
   DD + 'teava-pp-pentru-canalizare-interioara-fonoabsorbanta-50-mm-1-m/p/2034240',
   5 * 14.16 + 3 * 9 + 25 + 6 * 3.11 + 4 * 4 + 2 * 8 + 18 + 2 * 4 + 20, 'aceleași, la prețul mai mare',
   art='Canalizare PP: Ø50 ~5 m (lavoar, cadă), Ø40 ~3 m (mașina de spălat), Ø110 ~1 m (WC), coturi, ramificații',
   note='Lungimi din modelul băii: lavoar → coloană 1,7 m, sifon cadă → coloană 2,3 m, mașina de spălat ~2–3 m.')
bo('M052', 325 + 105.90, 'Leroy Merlin CL 51 15 kg 325 lei + Dedeman bandă CL 152 10 m 105,90 lei',
   'https://www.leroymerlin.ro/produse/emulsii-hidroizolante-si-bitum/505/hidroizolatie-flexibila-ceresit-cl51-wp-monocomponenta-15-kg/205595',
   2 * 355 + 2 * 105.90, 'Dedeman, 2 seturi dacă se hidroizolează și pereții', note='1 set pentru pardoseala băii și ridicările pe pereți; 2 seturi dacă se tratează toți pereții umezi.')
bo('M056', 40, 'sifon / racord scurgere mașină de spălat (estimat)', '', 80, 'estimat', st=E,
   art='Racord scurgere mașină de spălat (sifon cu clapetă), furtun', note='Țevile și robinetul mașinii sunt deja în M049–M051.')
# ================================================================ 7. incalzire
bo('M057', 450, 'Brico Depot 450 lei', 'https://www.bricodepot.ro/incalzire-si-climatizare/radiatoare/purmo-compact-22-600-x-1000.html', 455, 'Hornbach 455 lei')
bo('M061', 10 * 4.20 + 8 * 15.56 + 20, 'Dedeman: multistrat Firat 16 × 2 4,20 lei/m × 10; racord 16 × 2 15,56 lei × 8; izolație ~20 lei',
   DD + 'teava-pexal-firat-16-x-2-mm/p/2033742', 12 * 5.11 + 10 * 15.56 + 40, 'PEX-Al 5,11 lei/m',
   art='Racorduri radiatoare: multistrat 16 × 2 ~10 m, 8 racorduri, izolație', note='R1 și R2 lângă coloana C1 (~1 m fiecare, tur + retur); R3 în baie, racord de confirmat.')
# ================================================================ 8. placi si atelier
bo('M062', 366.42, 'Regency Company 366,42 lei (preț special; magazine în București, cu debitare)', 'https://regencycompany.ro/pal-egger-alb-premium-w1000-st9-18mm-2-80-x-2-07',
   389.95, 'Mambricolaj 389,95 lei', note='Metalcom / SmartPlusDeco au 302 lei, dar debitarea Metalcom e la Craiova; merită cerută ofertă cu transport.')
bo('M063', 366.42, 'Regency Company 366,42 lei (preț special)', 'https://regencycompany.ro/pal-melaminat-egger-gri-casmir-u702-st9-18mm-2-80-x-2-07', 409.95, 'Mambricolaj 409,95 lei')
bo('M064', 340, 'estimat (16 mm nu apare online)', 'https://www.mam-bricolaj.ro/cumpara/pal-melaminat-alb-premium-w1000-st9-2800x2070x18mm-7329', 390, 'estimat', st=E)
bo('M065', 126.65, 'Mathaus 126,65 lei', 'https://mathaus.ro/p/placa-hdf-lacuit-kronospan-alb-101-2800-x-2070-x-3-mm/000000000010103730', 149, 'Hornbach 149 lei')
bo('M068', 595.29, 'Regency 595,29 lei', 'https://regencycompany.ro/blat-bucatarie-egger-stejar-bardolino-natur-h1145-st10-38-x-600-x-4100', 659.95, 'Mambricolaj 659,95 lei')

# ---------------- debitare + cant calculate pe piese (planul de debitare + listele exacte ale dulapului din hol)
LAY = json.loads((ROOT / 'tmp/debitare_layout.json').read_text(encoding='utf-8'))
pcs = [p for m, r in LAY['B'].items() for s in r['sheets'] for p in s]
nsheets = sum(r['n'] for r in LAY['B'].values())
cut = sum(p['L'] + p['W'] for p in pcs) / 1000 + nsheets * 4.87          # m: două laturi pe piesă + tăiere margini placă
XS = {}
exec((ROOT / 'scripturi/excel_debitare.py').read_text(encoding='utf-8').split('# ---------------------------------------------------------------- gaurire')[0], XS)
c2 = c08 = e2 = e08 = 0.0
for p in XS['P']:                                                         # dulap hol: cantul exact din lista de debitare
    cod, den, mat, gr, Lmm, lmm, n, k1, k2, k3, k4 = p[:11]
    for val, ln in ((k1, Lmm), (k2, Lmm), (k3, lmm), (k4, lmm)):
        if val == '-':
            continue
        if val.startswith('2'):
            c2 += n * ln / 1000; e2 += n
        else:
            c08 += n * ln / 1000; e08 += n
for p in pcs:                                                             # restul: fronturi cant 2 mm pe 4 laturi, corpuri 1 latură
    if p['src'] == 'Dulap hol':
        continue
    Lm, Wm, mat, lab = p['L'] / 1000, p['W'] / 1000, p['mat'], p['label'].lower()
    if mat.startswith(('U702', 'H1145')):
        c2 += 2 * (Lm + Wm); e2 += 4
    elif mat.startswith('W1000 18'):
        if 'soclu' in lab:
            c08 += Lm; e08 += 1
        else:
            c2 += Lm; e2 += 1
    elif mat.startswith('W1000 16'):
        c08 += Lm; e08 += 1
# bucătăria (corpurile nu sunt încă desfăcute pe plăci): ~5 plăci W1000 + 1 U702 + 1 W16, la media celorlalte plăci
k_sheets = 7
cut += cut / nsheets * k_sheets
c2 += c2 / nsheets * k_sheets; c08 += c08 / nsheets * k_sheets; e2 += e2 / nsheets * k_sheets; e08 += e08 / nsheets * k_sheets
T = dict(deb=(1.79, 2.33), c2=(2.54 + 3.95, 5.00 + 4.07), c08=(1.80 + 0.80, 3.50))
deb_best = cut * T['deb'][0] + (c2 + e2 * 0.04) * T['c2'][0] + (c08 + e08 * 0.04) * T['c08'][0]
deb_alt = cut * T['deb'][1] + (c2 + e2 * 0.04) * T['c2'][1] + (c08 + e08 * 0.04) * T['c08'][1]
bo('M067', deb_best, 'Mambricolaj: debitare 1,79 lei/m, aplicare cant 22 × 2 2,54 lei/m + cant ABS W1000 3,95 lei/m', 'https://www.mam-bricolaj.ro/cumpara/aplicare-cant-22x2-3603',
   deb_alt, 'Darel: aplicare 5,00 lei/m + cant 4,07 lei/m', art='Debitare + cant ABS (calculat pe piese)', q=1, um='lot',
   note='%.0f m de tăiere, %.0f m cant 2 mm și %.0f m cant 0,8 mm, cu 4 cm pierdere la fiecare muchie; bucătăria estimată la media celorlalte plăci.' % (cut, c2, c08))
if 'M104' not in ROWS:
    i = [r['id'] for r in D['rows']].index('M070') + 1
    D['rows'].insert(i, dict(id='M104', cat=ROWS['M067']['cat'], art='Găurire CNC: locașuri balamale Ø35 (~100) și găuri confirmat (~400)', q=1, um='lot',
                             pmin=400, pmax=800, st=E, url='', note='Tarif de cerut la atelier odată cu debitarea; dulapul din hol are planul de găurire complet.',
                             price_basis='Estimat ~2–4 lei / locaș balama și ~0,5–1 leu / gaură (05.10.2026)', tmin=0, tmax=0))
    ROWS = {r['id']: r for r in D['rows']}

# ================================================================ 9. feronerie: seturi calculate pe bucăți, la cel mai mic preț găsit
U = {'h_apl': (6.77, 8.85, 'balama Blum 70T3550'), 'h_semi': (7.50, 12.99, 'balama Blum 70T3650'), 'plate': (0.85, 1.00, 'plăcuță 173L6100'),
     'tipon': (18.86, 26.78, 'TIP-ON 956.1004'), 'adapt': (4, 8, 'adaptor TIP-ON'), 'mov500': (124.25, 159.15, 'MOVENTO 760H5000S'),
     'mov350': (118.99, 139.00, 'MOVENTO 760H3500S'), 'tob': (85.83, 98.90, 'TIP-ON BLUMOTION S1'), 'sync': (17.24, 18.73, 'bară sincronizare'),
     'cupl': (6, 12, 'cleme cuplare'), 'conf': (0.0995, 0.254, 'confirmat'), 'capac': (0.0524, 0.10, 'capac'), 'cep': (0.05, 0.10, 'cep Ø8'),
     'coltar': (1.0, 2.0, 'colțar'), 'antitip': (5, 15, 'anti-răsturnare'), 'picior': (1.25, 1.87, 'picior H100'), 'clema': (2, 4, 'clemă soclu'),
     'pin': (0.10, 0.20, 'pin raft'), 'bara': (6.11, 21.78, 'bară ovală 1 m'), 'supbara': (0.76, 5.77, 'suport bară'), 'surub100': (8, 15, 'șuruburi ×100'),
     'cos': (188, 298, 'coș gunoi Ecomax 60'), 'picurator': (90, 150, 'picurător 80'), 'suspensie': (15, 40, 'agățătoare corp'),
     'cosrufe': (40, 80, 'coș rufe'), 'platforma': (300, 700, 'platformă mașină'), 'supraft': (3, 5, 'suport raft')}
VERIF = {'h_apl', 'h_semi', 'plate', 'tipon', 'mov500', 'mov350', 'tob', 'sync', 'conf', 'capac', 'picior', 'bara', 'supbara', 'cos'}
SETS = {
    'M071': [('h_apl', 6), ('h_semi', 18), ('plate', 24), ('tipon', 8), ('adapt', 8), ('mov500', 2), ('tob', 2), ('sync', 2), ('cupl', 2), ('conf', 50), ('capac', 50),
             ('cep', 4), ('coltar', 10), ('antitip', 3), ('picior', 10), ('clema', 5), ('pin', 16), ('bara', 2), ('supbara', 4), ('surub100', 4)],
    'M072': [('h_apl', 6), ('h_semi', 24), ('plate', 30), ('tipon', 10), ('adapt', 10), ('mov500', 6), ('cupl', 6), ('conf', 110), ('capac', 60), ('coltar', 8),
             ('antitip', 3), ('picior', 12), ('clema', 6), ('pin', 44), ('bara', 3), ('supbara', 6), ('surub100', 4)],
    'M074': [('h_apl', 10), ('h_semi', 14), ('plate', 24), ('tipon', 10), ('adapt', 10), ('mov500', 3), ('tob', 3), ('sync', 3), ('cupl', 3), ('conf', 80), ('capac', 40),
             ('picior', 16), ('clema', 8), ('cos', 1), ('picurator', 1), ('suspensie', 10), ('surub100', 4)],
    'M075': [('h_apl', 4), ('h_semi', 6), ('plate', 10), ('tipon', 4), ('adapt', 4), ('mov500', 2), ('tob', 1), ('cupl', 2), ('cosrufe', 2), ('platforma', 1),
             ('conf', 40), ('capac', 30), ('coltar', 4), ('antitip', 2), ('picior', 6), ('surub100', 2)],
    'M076': [('mov350', 8), ('tob', 8), ('cupl', 8), ('h_apl', 6), ('h_semi', 4), ('plate', 10), ('tipon', 5), ('adapt', 5), ('suspensie', 7), ('supraft', 24),
             ('conf', 80), ('capac', 40), ('surub100', 2)]}
SETART = {'M071': 'Dulap hol: feronerie Blum completă (calculată pe bucăți)', 'M072': 'Dressing: 30 balamale, 10 TIP-ON, 6 sertare MOVENTO, bare, accesorii',
          'M074': 'Bucătărie: 24 balamale, 10 TIP-ON, 3 sertare MOVENTO cu TIP-ON BLUMOTION, coș gunoi, picurător, picioare',
          'M075': 'Spălătorie: 10 balamale, 4 TIP-ON, sertar coșuri + poliță extractibilă MOVENTO, platformă mașină',
          'M076': 'Noptiere, comode TV, machiaj, dulap WC, debara: 8 sertare MOVENTO + TIP-ON BLUMOTION, 10 balamale, suporți'}
for sid, items in SETS.items():
    b = sum(q * U[k][0] for k, q in items); a = sum(q * U[k][1] for k, q in items)
    share = sum(q * U[k][0] for k, q in items if k in VERIF) / b
    det = '; '.join('%g × %s %s' % (q, U[k][2], ('%.2f' % U[k][0]).replace('.', ',')) for k, q in items)
    bo(sid, b, 'cel mai mic preț pe bucată (Feromob, Triton, Darel, Regency, Feroshop, Mambricolaj)',
       'https://www.feromob.ro/product/balama-blum-clip-top-aplicata-110-fara-arc70t3550-tl/', a, 'Regency / Schachermayer', st=V if share > 0.8 else E,
       art=SETART[sid], note='Cantități și prețuri: ' + det + '.')
bo('M073', 729.99, 'Feroshop 729,99 lei', 'https://www.feroshop.ro/feronerie/feronerie-functionala/organizare/dressing/ironfix-sertar-masa-de-calcat-56860710-40384.html', 1019, 'Luxfer 1.019 lei')
# ================================================================ 10. bucatarie
bo('M078', 469, 'Demax 469 lei', 'https://baterii-bucatarie.compari.ro/grohe/30463000-p681805632/', 519, 'compari.ro 519 lei')
bo('M079', 396.99, 'eMAG: Heinner HTCH-MD600GBK, 630 m³/h, 396,99 lei', 'https://www.emag.ro/search/hota+60cm', 1446.23, 'Electrolux LFP616X 1.446,23 lei')
# ================================================================ 12. mobilier gata facut si textile
bo('M092', 1250, 'JYSK ILDVED 1.250 lei (192 cm, ladă, pat 114 × 192)', 'https://jysk.ro/sufragerie/canapele-extensibile/canapea-extensibila-ildved-bej',
   1999, 'IKEA FRIHETEN 1.999 lei (~230 cm, mai aproape de randare)', art='Canapea extensibilă cu ladă (JYSK ILDVED / IKEA FRIHETEN)')
bo('M093', 1000, 'JYSK cadru pat HALD 160 × 200 stejar, cu depozitare, 1.000 lei', 'https://jysk.ro/dormitor/cadre-de-pat/cadru-pat-hald-160x200-stejar',
   1699, 'IKEA GLADSTAD tapițat, 4 cutii, 1.699 lei', art='Pat 160 × 200 cu depozitare (JYSK HALD / IKEA GLADSTAD)')
bo('M096', 299, 'IKEA FLINTAN 299 lei', 'https://www.ikea.com/ro/ro/p/flintan-scaun-rotativ-negru-10489028/', 399, 'IKEA FLINTAN cu brațe 399 lei')
bo('M098', 239 + 253.41, 'IKEA MORUM 160 × 230 239 lei (IKEA Family) + puf 253,41 lei (eMAG)', 'https://www.ikea.com/ro/ro/p/morum-covor-tesatura-plata-int-ext-bej-20203563/',
   449 + 429.99, 'IKEA LOHALS 449 lei + puf 429,99 lei')
bo('M099', 2 * 209 + 2 * 145 + 200, 'IKEA MAJGULL opace 2 buc 209 lei × 2 + șine VIDGA 145 lei × 2 + voal ~200 lei', 'https://www.ikea.com/ro/ro/p/majgull-draperii-opace-2buc-verde-inchis-rjns-30586033/',
   2 * 249 + 2 * 145 + 400, 'MAJGULL gri 249 lei + voal ~400 lei', st=E)

# ================================================================ etape de cheltuieli
STAGES = [('E1', 'Instalații brute și comenzi lungi', 1), ('E2', 'Șape, hidroizolații, măști de rigips', 2), ('E3', 'Finisaje: microciment, epoxid, parchet, vopsea', 3),
          ('E4', 'Mobilier la comandă: plăci, atelier, feronerie, LED', 3), ('E5', 'Aparataj, protecții, iluminat, obiecte sanitare', 5),
          ('E6', 'Bucătărie, electrocasnice și TV (Black Friday)', 5), ('E7', 'Mobilier gata făcut și textile', 6), ('E8', 'Transport și livrări (pe parcurs)', None)]
ET = {}
for s, ids in {'E1': 'M018 M019 M020 M023 M026 M028 M030 M032 M033 M040 M041 M043 M049 M050 M051 M055 M056 M057 M058 M059 M060 M061 M102',
               'E2': 'M001 M002 M009 M052 M053', 'E3': 'M003 M004 M005 M006 M007 M008 M010 M011 M012 M013 M014 M015 M016 M017',
               'E4': 'M036 M037 M038 M039 M062 M063 M064 M065 M066 M067 M068 M069 M070 M071 M072 M073 M074 M075 M076 M104',
               'E5': 'M021 M022 M024 M025 M027 M029 M031 M034 M035 M042 M044 M045 M046 M047 M048 M054',
               'E6': 'M077 M078 M079 M080 M083 M084 M085 M086 M087 M088 M089 M090 M091', 'E7': 'M081 M082 M092 M093 M094 M095 M096 M097 M098 M099 M100 M101',
               'E8': 'M103'}.items():
    for i in ids.split():
        ET[i] = s
missing = [r['id'] for r in D['rows'] if r['id'] not in ET]
assert not missing, missing
for r in D['rows']:
    r['etapa'] = ET[r['id']]
    r['tmin'] = round(r['q'] * r['pmin'], 2)
    r['tmax'] = round(r['q'] * r['pmax'], 2)
    r.setdefault('magazin', ''); r.setdefault('magazin_alt', '')
cats = list(dict.fromkeys(r['cat'] for r in D['rows']))
D['summary'] = [dict(cat=c, tmin=round(sum(r['tmin'] for r in D['rows'] if r['cat'] == c), 2), tmax=round(sum(r['tmax'] for r in D['rows'] if r['cat'] == c), 2),
                     n=sum(r['cat'] == c for r in D['rows']), ver=sum(r['cat'] == c and r['st'] == V for r in D['rows'])) for c in cats]
D['old_total'] = [json.loads(BASE.read_text(encoding='utf-8'))[k] for k in ('tot_min', 'tot_max')]
D['tot_min'] = round(sum(s['tmin'] for s in D['summary']), 2)
D['tot_max'] = round(sum(s['tmax'] for s in D['summary']), 2)
D['ver_sum'] = round(sum(r['tmin'] for r in D['rows'] if r['st'] == V), 2)
D['revision'] = 'R03'
WEEKS = [('S0', datetime.date(2026, 10, 5))] + [('S%d' % k, datetime.date(2026, 10, 5) + datetime.timedelta(weeks=k)) for k in range(1, 9)]
st_rows = []
for code, name, wk in STAGES:
    rs = [r for r in D['rows'] if r['etapa'] == code]
    st_rows.append(dict(cod=code, nume=name, sapt=wk, n=len(rs), tmin=round(sum(r['tmin'] for r in rs), 2), tmax=round(sum(r['tmax'] for r in rs), 2)))
# săptămâna în care se cumpără fiecare articol: implicit cea a etapei, cu excepțiile de mai jos (cumpărat chiar înainte de folosire)
WK_OVR = {'M053': 2, 'M011': 2, 'M012': 2, 'M013': 4, 'M014': 4, 'M015': 2, 'M006': 4, 'M010': 4, 'M007': 4,
          'M003': 5, 'M004': 5, 'M005': 5, 'M008': 5, 'M016': 5, 'M017': 5, 'M036': 6, 'M037': 6, 'M038': 6, 'M039': 6, 'M001': 3, 'M002': 3, 'M009': 3, 'M052': 3}
SW = {c: w for c, _, w in STAGES}
for r in D['rows']:
    r['sapt'] = WK_OVR.get(r['id'], SW[r['etapa']])
for srow in st_rows:
    ws_ = sorted({r['sapt'] for r in D['rows'] if r['etapa'] == srow['cod'] and r['sapt']})
    srow['sapt_txt'] = ('S%d' % ws_[0] if len(ws_) == 1 else 'S%d–S%d' % (ws_[0], ws_[-1])) if ws_ else 'S1–S7'
cash = collections.OrderedDict((w, [0.0, 0.0]) for w, _ in WEEKS)
what = collections.defaultdict(set)
for r in D['rows']:
    if r['sapt'] is None:                                                 # transportul se împarte pe săptămânile 1–7
        for k in range(1, 8):
            cash['S%d' % k][0] += r['tmin'] / 7; cash['S%d' % k][1] += r['tmax'] / 7; what['S%d' % k].add('transport')
    else:
        cash['S%d' % r['sapt']][0] += r['tmin']; cash['S%d' % r['sapt']][1] += r['tmax']; what['S%d' % r['sapt']].add(r['etapa'])
D['etape'] = st_rows
D['cashflow'] = [dict(sapt=w, data=d.strftime('%d.%m.%Y'), tmin=round(cash[w][0], 2), tmax=round(cash[w][1], 2), ce=', '.join(sorted(what[w]))) for w, d in WEEKS]
nv = sum(r['st'] == V for r in D['rows'])
D['notes'] = [n for n in D['notes'] if not n.startswith(('R02:', 'R03:'))] + [
    'R03: pentru fiecare articol, prețul de bază este cea mai bună ofertă găsită online pe 05.10.2026 (magazin și link), iar alternativa este prețul la un magazin comod. '
    '%d din %d linii au prețul verificat. Seturile de feronerie Blum și debitarea sunt calculate pe bucăți și pe metri, din planul de debitare.' % (nv, len(D['rows']))]

# ================================================================ calendarul lucrărilor
d = lambda s: datetime.date(2026, int(s[3:5]), int(s[:2]))
TL = [('Pregătire', 'Releveu și verificări: tavan hol ≥ 2480 mm, coloane apă / canal / termic, tablou, gaz', '05.10', '09.10', 'beneficiar + executanți'),
      ('Pregătire', 'Oferte executanți; electrician ANRE; acord asociație și programare golire coloană termică', '05.10', '09.10', 'beneficiar'),
      ('Pregătire', 'Comenzi etapa E1: radiatoare, Geberit, cadă, materiale de instalații', '07.10', '12.10', 'beneficiar'),
      ('Instalații', 'Demontări: radiatoare de fontă, obiecte sanitare, finisaje vechi; protecții', '12.10', '13.10', 'echipă'),
      ('Instalații', 'Termic: racorduri noi, radiatoare R1 / R2 / R3, probă — înainte de pornirea căldurii', '13.10', '15.10', 'instalator'),
      ('Instalații', 'Sanitar brut: alimentare PPR, canalizare PP, stelaj Geberit, cadă, probă de presiune', '13.10', '17.10', 'instalator'),
      ('Instalații', 'Electric brut: trasare, tuburi, doze, cabluri 230 V și date, tablou', '13.10', '21.10', 'electrician ANRE'),
      ('Pereți', 'Măști de rigips în baie (Geberit, nișe, trape), reparații, închidere șanțuri', '20.10', '24.10', 'echipă'),
      ('Pereți', 'Glet pe pereți și tavane, 2 straturi, șlefuire', '22.10', '30.10', 'echipă'),
      ('Pardoseli', 'Amorsă + șapă autonivelantă pe toată suprafața (apoi 7 zile de uscare)', '26.10', '26.10', 'echipă'),
      ('Mobilier', 'Comandă debitare, cant, găurire și feronerie (Excelurile pentru atelier)', '26.10', '26.10', 'beneficiar'),
      ('Pardoseli', 'Hidroizolație baie: CL 51, 2 straturi + bandă CL 152', '28.10', '29.10', 'echipă'),
      ('Pardoseli', 'Microciment baie, hol, vestibul (straturi + sigilare)', '02.11', '07.11', 'aplicator'),
      ('Pereți', 'Vopsea, primul strat: camera de zi, dormitor, bucătărie', '02.11', '04.11', 'echipă'),
      ('Pardoseli', 'Rășină epoxidică: bucătărie, debara, balcon (+72 h întărire)', '05.11', '08.11', 'echipă'),
      ('Mobilier', 'Black Friday: electrocasnice și televizoare (6 sau 13.11)', '06.11', '13.11', 'beneficiar'),
      ('Mobilier', 'Livrarea pieselor debitate (~10 zile lucrătoare de la comandă)', '09.11', '10.11', 'atelier'),
      ('Pardoseli', 'Parchet + subparchet, plinte, praguri', '09.11', '11.11', 'echipă'),
      ('Pereți', 'Vopsea finală, baghete perete TV, riflaje dormitor', '11.11', '13.11', 'echipă'),
      ('Finisare', 'Aparataj, plafoniere, aplice, protecții în tablou, măsurători PRAM', '12.11', '14.11', 'electrician ANRE'),
      ('Finisare', 'Obiecte sanitare: WC, lavoar, baterii, paravan', '12.11', '13.11', 'instalator'),
      ('Mobilier', 'Montaj dulap hol, dressing, spălătorie', '16.11', '19.11', 'montator'),
      ('Mobilier', 'Montaj bucătărie, blat, chiuvetă, plită (firmă autorizată gaz), hotă, cuptor, MSV', '18.11', '21.11', 'montator + gaz'),
      ('Mobilier', 'Noptiere, comode TV, machiaj, dulap WC, rafturi debara, LED', '19.11', '21.11', 'montator'),
      ('Finisare', 'Mobilier gata făcut, draperii, rulouri; TV pe perete', '23.11', '25.11', 'beneficiar'),
      ('Finisare', 'Oglinzi lipite, reglaje uși, curățenie finală', '24.11', '27.11', 'echipă'),
      ('Finisare', 'Rezervă pentru întârzieri; fotografii pentru anunț', '30.11', '04.12', 'beneficiar')]
TLJ = [dict(faza=f, act=a, start=d(s).isoformat(), end=d(e).isoformat(), cine=c) for f, a, s, e, c in TL]
(ROOT / 'tmp/timeline.json').write_text(json.dumps(dict(weeks=[(w, x.isoformat()) for w, x in WEEKS], tasks=TLJ), ensure_ascii=False, indent=1), encoding='utf-8')
SRC.write_text(json.dumps(D, ensure_ascii=False, indent=2), encoding='utf-8')

# ================================================================ Excel R03
THIN = Side(style='thin', color='C8CCD2'); BR = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill('solid', fgColor='147B73'); VER = PatternFill('solid', fgColor='E7F3F1'); PANEL = PatternFill('solid', fgColor='E7F3F1')
BAR = PatternFill('solid', fgColor='147B73'); BAR2 = PatternFill('solid', fgColor='9FD3CC')


def head(ws, row, labels):
    for ci, h in enumerate(labels, start=1):
        c = ws.cell(row=row, column=ci, value=h); c.font = Font(bold=True, color='FFFFFF'); c.fill = HEAD; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BR


wb = Workbook(); ws = wb.active; ws.title = 'Rezumat'
ws['A1'] = 'Deviz materiale R03 – Apartament 184, Bloc L14 (fără manoperă)'; ws['A1'].font = Font(bold=True, size=14, color='147B73')
ws['A2'] = '„Cea mai bună ofertă” = cel mai mic preț verificat online pe 05.10.2026, cu link; „Alternativă” = prețul la un magazin comod. Lei cu TVA.'
head(ws, 4, ['Categorie', 'Cea mai bună ofertă (lei)', 'Alternativă (lei)', 'Linii', 'Verificate'])
r0 = 5
for s in D['summary']:
    ws.append([s['cat'], s['tmin'], s['tmax'], s['n'], s['ver']])
r1 = ws.max_row
ws.append(['Subtotal', '=SUM(B%d:B%d)' % (r0, r1), '=SUM(C%d:C%d)' % (r0, r1)])
sub = ws.max_row
ws.append(['Rezervă neprevăzute', '=B%d*$E$%d' % (sub, sub + 1), '=C%d*$E$%d' % (sub, sub + 1), 'procent:', D['rez']])
ws.append(['TOTAL MATERIALE', '=B%d+B%d' % (sub, sub + 1), '=C%d+C%d' % (sub, sub + 1)])
for rr in range(sub, sub + 3):
    ws.cell(row=rr, column=1).font = Font(bold=True)
ws.cell(row=sub + 1, column=5).number_format = '0%'
ws.append([])
head(ws, ws.max_row + 1, ['Etapă de cheltuieli', 'Cea mai bună ofertă (lei)', 'Alternativă (lei)', 'Linii', 'Săptămâna plății'])
for s in st_rows:
    ws.append(['%s – %s' % (s['cod'], s['nume']), s['tmin'], s['tmax'], s['n'], s['sapt_txt']])
for row in ws.iter_rows(min_row=5, max_row=ws.max_row):
    for c in row:
        if c.value is not None:
            c.border = BR
        if c.column in (2, 3):
            c.number_format = '#,##0'
for col, w in zip('ABCDE', (62, 22, 18, 8, 18)):
    ws.column_dimensions[col].width = w

ws2 = wb.create_sheet('Deviz detaliat')
H0 = 5; n = H0 + len(D['rows'])
hdr = ['ID', 'Etapă', 'Categorie', 'Articol', 'Cant.', 'UM', 'Cea mai bună ofertă', 'Alternativă', 'Total ofertă', 'Total alternativă', 'Status', 'Link', 'Unde e oferta', 'Alternativa', 'Bază cantitate / notă']
ws2['A1'] = 'TOTAL SELECȚIE – se recalculează după filtrele din antet (Etapă, Categorie, Status…)'; ws2['A1'].font = Font(bold=True, size=12, color='147B73')
ws2['I1'], ws2['J1'], ws2['K1'] = 'Oferte', 'Alternativă', 'Mediu'
for rr, lab, f in ((2, 'Fără rezervă', '=SUBTOTAL(109,{c}%d:{c}%d)' % (H0 + 1, n)), (3, 'Cu rezervă %d %%' % (D['rez'] * 100), '={c}2*%s' % (1 + D['rez']))):
    ws2.cell(row=rr, column=4, value=lab).font = Font(bold=True)
    for col in 'IJ':
        c = ws2['%s%d' % (col, rr)]; c.value = f.format(c=col); c.number_format = '#,##0'; c.font = Font(bold=True, size=12); c.fill = PANEL; c.border = BR
    ws2['K%d' % rr] = '=(I%d+J%d)/2' % (rr, rr); ws2['K%d' % rr].number_format = '#,##0'; ws2['K%d' % rr].fill = PANEL
ws2['D4'] = 'Linii afișate:'; ws2['E4'] = '=SUBTOTAL(103,A%d:A%d)' % (H0 + 1, n); ws2['D4'].font = Font(bold=True)
head(ws2, H0, hdr)
for i, r in enumerate(D['rows'], start=H0 + 1):
    vals = [r['id'], r['etapa'], r['cat'], r['art'], r['q'], r['um'], r['pmin'], r['pmax'], '=E%d*G%d' % (i, i), '=E%d*H%d' % (i, i), r['st'], r['url'] or '',
            r.get('magazin') or r.get('price_basis', ''), r.get('magazin_alt', ''), r['note']]
    for ci, v in enumerate(vals, start=1):
        c = ws2.cell(row=i, column=ci, value=v); c.border = BR; c.alignment = Alignment(wrap_text=True, vertical='top')
        if ci in (7, 8, 9, 10):
            c.number_format = '#,##0.00'
        if r['st'] == V:
            c.fill = VER
    if r['url']:
        c = ws2.cell(row=i, column=12); c.hyperlink = r['url']; c.value = 'deschide'; c.font = Font(color='277DBA', underline='single')
ws2.auto_filter.ref = 'A%d:O%d' % (H0, n)
for col, w in zip('ABCDEFGHIJKLMNO', (7, 7, 22, 50, 7, 7, 12, 12, 12, 12, 10, 10, 46, 30, 46)):
    ws2.column_dimensions[col].width = w
ws2.freeze_panes = 'E%d' % (H0 + 1)

ws3 = wb.create_sheet('Etape și plăți')
ws3['A1'] = 'Plăți pe săptămâni (materiale, fără manoperă)'; ws3['A1'].font = Font(bold=True, size=13, color='147B73')
head(ws3, 3, ['Săptămâna', 'Luni', 'Cea mai bună ofertă (lei)', 'Alternativă (lei)', 'Cumulat ofertă', 'Cumulat alternativă', 'Ce se plătește'])
for k, cf in enumerate(D['cashflow'], start=4):
    ws3.append([cf['sapt'], cf['data'], cf['tmin'], cf['tmax'], '=SUM($C$4:C%d)' % k, '=SUM($D$4:D%d)' % k, cf['ce']])
for row in ws3.iter_rows(min_row=4, max_row=ws3.max_row):
    for c in row:
        c.border = BR
        if c.column in (3, 4, 5, 6):
            c.number_format = '#,##0'
for col, w in zip('ABCDEFG', (11, 12, 22, 18, 16, 18, 40)):
    ws3.column_dimensions[col].width = w

ws4 = wb.create_sheet('Calendar lucrări')
ws4['A1'] = 'Calendarul lucrărilor – start luni 12.10.2026, final ~27.11.2026, rezervă până la 04.12'; ws4['A1'].font = Font(bold=True, size=13, color='147B73')
labels = ['Fază', 'Activitate', 'Început', 'Sfârșit', 'Cine'] + ['%s\n%s' % (w, x.strftime('%d.%m')) for w, x in WEEKS]
head(ws4, 3, labels)
for i, t in enumerate(TLJ, start=4):
    s0, e0 = datetime.date.fromisoformat(t['start']), datetime.date.fromisoformat(t['end'])
    for ci, v in enumerate([t['faza'], t['act'], s0.strftime('%d.%m'), e0.strftime('%d.%m'), t['cine']], start=1):
        c = ws4.cell(row=i, column=ci, value=v); c.border = BR; c.alignment = Alignment(wrap_text=True, vertical='top')
    for k, (w, ws_) in enumerate(WEEKS):
        we = ws_ + datetime.timedelta(days=6)
        c = ws4.cell(row=i, column=6 + k); c.border = BR
        if s0 <= we and e0 >= ws_:
            c.fill = BAR if t['faza'] != 'Pregătire' else BAR2
ws4.column_dimensions['A'].width = 12; ws4.column_dimensions['B'].width = 70; ws4.column_dimensions['C'].width = 8; ws4.column_dimensions['D'].width = 8; ws4.column_dimensions['E'].width = 18
for k in range(len(WEEKS)):
    ws4.column_dimensions[chr(ord('F') + k)].width = 7
ws4.row_dimensions[3].height = 32
ws4.freeze_panes = 'F4'

ws5 = wb.create_sheet('Opționale')
head(ws5, 1, ['Articol', 'Cant.', 'UM', 'Preț min', 'Preț max', 'Notă'])
for o in D['opt']:
    ws5.append(list(o))
out = ROOT / 'Deviz_materiale_ap184_R03.xlsx'
wb.save(out)
print('R03', collections.Counter(r['st'] for r in D['rows']), 'linii', len(D['rows']), 'cu link', sum(bool(r['url']) for r in D['rows']))
print('oferte %.0f | alternativă %.0f | cu rezervă %.0f – %.0f | R02 %.0f – %.0f' % (D['tot_min'], D['tot_max'], D['tot_min'] * 1.07, D['tot_max'] * 1.07, *D['old_total']))
print('debitare: %.0f m taiere, %.0f m cant 2 mm, %.0f m cant 0,8 → %.0f – %.0f lei' % (cut, c2, c08, deb_best, deb_alt))
for s in D['summary']:
    print('%-45s %9.0f %9.0f  %d/%d' % (s['cat'], s['tmin'], s['tmax'], s['ver'], s['n']))
for s in st_rows:
    print(s['cod'], s['nume'][:40], round(s['tmin']), round(s['tmax']))
for cf in D['cashflow']:
    print(cf['sapt'], cf['data'], round(cf['tmin']), round(cf['tmax']))
